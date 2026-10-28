"""Build the forbidden-wording report (HTML + XLSX + zip) from a crawl.mjs run.

  python3 report.py --crawl <crawl out dir> --rules <rules.json> --site "Win777" --out <report dir>

Output in <report dir>:
  <Site>_Forbidden_Wording_Report.html   one section per page, findings table, screenshots
  <Site>_Forbidden_Wording_Findings.xlsx  one row per finding with a Done column, plus Pages and Exempt sheets
  screenshots/                            only the screenshots the report references
and <report dir>.zip next to it, which is what gets sent to the person fixing the copy.
"""
import argparse, datetime, html, json, os, re, shutil
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument('--crawl', required=True)
ap.add_argument('--rules', required=True)
ap.add_argument('--site', required=True, help='display name, e.g. Win777')
ap.add_argument('--out', required=True)
ap.add_argument('--date', default=datetime.date.today().strftime('%-d %b %Y'))
a = ap.parse_args()

RULESET = json.load(open(a.rules))
RULES = RULESET['rules']
BY_ID = {r['id']: r for r in RULES}
ORDER = [r['id'] for r in sorted(RULES, key=lambda r: {'Critical': 0, 'High': 1, 'Medium': 2, 'Low': 3}.get(r['severity'], 2))]
SHOT_RULES = {r['id'] for r in RULES if r.get('screenshot', True)}
EXEMPT_FIELDS = set(RULESET.get('slugExemptFields', []))
FILE_STEM = re.sub(r'[^A-Za-z0-9]+', '_', a.site).strip('_')
os.makedirs(os.path.join(a.out, 'screenshots'), exist_ok=True)

# last record per requested URL wins (retries), then collapse redirects to the same final URL
recs = {}
for l in open(os.path.join(a.crawl, 'results.jsonl')):
    if l.strip():
        r = json.loads(l); recs[r['req']] = r
seen, pages = set(), []
for r in recs.values():
    key = r.get('url') or r['req']
    if key not in seen:
        seen.add(key); pages.append(r)

broken = [p for p in pages if not p.get('error') and (re.search(r'not found|^404\b', p.get('title') or '', re.I) or (p.get('h1') or '').strip() == '404')]
errors = [p for p in pages if p.get('error')]
pages = [p for p in pages if p not in broken and p not in errors]

# header/footer text counts as site-wide only when it repeats on 3+ pages (blog hero bars often carry "header" classes)
rep = Counter()
for p in pages:
    for k in {(h['region'], h['rule'], h['context']) for h in p['hits'] if h['region'] != 'main'}:
        rep[k] += 1
for p in pages:
    for h in p['hits']:
        if h['region'] != 'main' and rep[(h['region'], h['rule'], h['context'])] < 3:
            h['region'] = 'main'


def slug_exempt(path, rule_id, field):
    tok = BY_ID[rule_id].get('slugToken')
    return bool(tok) and field in EXEMPT_FIELDS and tok in path.lower()


def hl(text, term):
    t = html.escape(text)
    if not term:
        return t
    return re.sub(re.escape(html.escape(term)), lambda m: f'<mark>{m.group(0)}</mark>', t, flags=re.I)


def ctx(text, term, width=220):
    i = text.lower().find(term.lower()) if term else -1
    if i < 0 or len(text) <= width * 2:
        return text
    lo, hi = max(0, i - width), min(len(text), i + len(term) + width)
    return ('…' if lo else '') + text[lo:hi] + ('…' if hi < len(text) else '')


rows, shared = [], defaultdict(lambda: {'pages': set(), 'hit': None})
for p in pages:
    url, path = p.get('url', p['req']), p.get('path') or '/'
    for h in p['head']:
        rows.append(dict(url=url, where=h['field'], rule=h['rule'], match=h['match'], heading='(page <head>)', context=h['text'],
                         visible='n/a', shot='', exempt=slug_exempt(path, h['rule'], h['field'])))
    for h in p['hits']:
        if h['region'] != 'main':
            s = shared[(h['region'], h['rule'], h['context'])]
            s['pages'].add(url)
            if h.get('shot'):
                s['hit'] = h
            continue
        is_h1 = h['inHeading'] and h['headingTag'] == 'H1' and h['match'].lower() in h['heading'].lower()
        where = 'H1' if is_h1 else ('Heading' if h['inHeading'] else 'Body')
        rows.append(dict(url=url, where=where, rule=h['rule'], match=h['match'], heading=h['heading'], context=h['context'],
                         visible='Visible' if h['visible'] else 'Collapsed (FAQ / accordion)', shot=h.get('shot', ''),
                         exempt=slug_exempt(path, h['rule'], where)))

used = {r['shot'] for r in rows if r['shot']} | {v['hit']['shot'] for v in shared.values() if v['hit'] and v['hit'].get('shot')}
for s in used:
    shutil.copy(os.path.join(a.crawl, s), os.path.join(a.out, 'screenshots', os.path.basename(s)))

actionable = [r for r in rows if not r['exempt']]
by_rule = Counter(r['rule'] for r in actionable)
page_rows = defaultdict(list)
for r in actionable:
    page_rows[r['url']].append(r)
top = ORDER[0]
word_pages = {u for u, rs in page_rows.items() if any(r['rule'] in SHOT_RULES for r in rs)}


def page_key(u):
    rs = page_rows[u]
    return (-sum(r['rule'] == top for r in rs), -sum(r['rule'] in SHOT_RULES for r in rs), u)


def lbl(rid):
    return BY_ID[rid]['label']


# ---------------- XLSX ----------------
HDR_FILL = PatternFill('solid', fgColor='7A0010')


def style_sheet(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = HDR_FILL
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions


wb = Workbook()
ws = wb.active; ws.title = 'Findings'
ws.append(['#', 'Page URL', 'Severity', 'Issue', 'Matched text', 'Where on page', 'Section (nearest heading)', 'Visibility',
           'Sentence / text block', 'Fix', 'Screenshot file', 'Done'])
n = 0
for u in sorted(page_rows, key=page_key):
    for r in sorted(page_rows[u], key=lambda r: ORDER.index(r['rule'])):
        n += 1
        ws.append([n, r['url'], BY_ID[r['rule']]['severity'], lbl(r['rule']), r['match'], r['where'], r['heading'], r['visible'],
                   ctx(r['context'], r['match']), BY_ID[r['rule']]['fix'], os.path.basename(r['shot']) if r['shot'] else '', ''])
for (region, rid, text), v in shared.items():
    n += 1
    ws.append([n, f"Site-wide {region} ({len(v['pages'])} pages)", BY_ID[rid]['severity'], lbl(rid), '', region.capitalize(), '',
               'Visible', text, BY_ID[rid]['fix'], os.path.basename(v['hit']['shot']) if v['hit'] else '', ''])
style_sheet(ws, [5, 48, 10, 26, 22, 14, 40, 16, 80, 50, 34, 8])

ws2 = wb.create_sheet('Pages')
ws2.append(['Page URL', 'In sitemap'] + [lbl(r) for r in ORDER] + ['Total (screenshot rules)', 'Screenshots'])
for p in pages:
    u = p.get('url', p['req']); c = Counter(r['rule'] for r in page_rows.get(u, []))
    ws2.append([u, 'Yes' if p.get('inSitemap') else 'No (linked only)'] + [c[r] for r in ORDER]
               + [sum(v for k, v in c.items() if k in SHOT_RULES), len(p.get('shots', []))])
style_sheet(ws2, [60, 16] + [12] * len(ORDER) + [14, 12])

ws3 = wb.create_sheet('Exempt')
ws3.append(['Page URL', 'Field', 'Issue', 'Text', 'Why exempt'])
for r in rows:
    if r['exempt']:
        ws3.append([r['url'], r['where'], lbl(r['rule']), r['context'], RULESET.get('slugExemptNote', 'Slug contains the avoided word.')])
style_sheet(ws3, [60, 16, 18, 70, 50])
wb.save(os.path.join(a.out, f'{FILE_STEM}_Forbidden_Wording_Findings.xlsx'))

# ---------------- HTML ----------------
E = html.escape


def badge(rid, count):
    sev = BY_ID[rid]['severity']
    cls = {'Critical': 'crit', 'High': 'hi', 'Medium': 'hi'}.get(sev, 'lo')
    return f'<span class="b {cls}">{E(lbl(rid))} × {count}</span>'


def shot_link(s):
    return f'<a href="screenshots/{E(os.path.basename(s))}" target="_blank">{E(os.path.basename(s))}</a>' if s else '—'


kpis = ''.join(f'<div class="kpi"><b>{by_rule.get(r, 0)}</b><span>{E(lbl(r))}</span></div>' for r in ORDER)
toc, sections = '', ''
for u in sorted(page_rows, key=page_key):
    rs = sorted(page_rows[u], key=lambda r: ORDER.index(r['rule']))
    main_rows = [r for r in rs if r['rule'] in SHOT_RULES]
    minor = [r for r in rs if r['rule'] not in SHOT_RULES]
    badges = ' '.join(badge(rid, c) for rid, c in sorted(Counter(r['rule'] for r in rs).items(), key=lambda x: ORDER.index(x[0])))
    pid = re.sub(r'[^a-z0-9]+', '-', u.lower())
    toc += f'<li><a href="#{pid}">{E(re.sub(r"^https?://[^/]+", "", u) or "/")}</a> {badges}</li>'
    trs = ''.join(
        f'<tr><td>{i}</td><td>{badge(r["rule"], 1).replace(" × 1", "")}</td><td>{E(r["where"])}<br><small>{E(r["visible"])}</small></td>'
        f'<td>{E(r["heading"])}</td><td class="q">{hl(ctx(r["context"], r["match"]), r["match"])}</td>'
        f'<td>{E(BY_ID[r["rule"]]["fix"])}</td><td>{shot_link(r["shot"])}</td></tr>' for i, r in enumerate(main_rows, 1))
    table = (f'<div class="tw"><table><thead><tr><th>#</th><th>Issue</th><th>Where</th><th>Section (nearest heading)</th><th>Text on the page</th>'
             f'<th>Fix</th><th>Screenshot</th></tr></thead><tbody>{trs}</tbody></table></div>') if main_rows else '<p class="muted">Only low-severity items on this page.</p>'
    fold = ''
    for rid, grp in sorted(defaultdict(list, {k: [r for r in minor if r['rule'] == k] for k in {r['rule'] for r in minor}}).items(), key=lambda x: ORDER.index(x[0])):
        items = ''.join(f'<li><small>{E(r["where"])} · {E(r["heading"][:90])}</small><br>{hl(ctx(r["context"], r["match"], 110), r["match"])}</li>' for r in grp)
        fold += f'<details><summary>{len(grp)} × {E(lbl(rid))}: {E(BY_ID[rid]["fix"])}</summary><ol class="dl">{items}</ol></details>'
    gal = ''.join(f'<a href="screenshots/{E(os.path.basename(s))}" target="_blank"><img loading="lazy" src="screenshots/{E(os.path.basename(s))}" alt="Screenshot {E(os.path.basename(s))}"><span>{E(os.path.basename(s))}</span></a>'
                  for s in sorted({r['shot'] for r in main_rows if r['shot']}))
    sections += f'<section id="{pid}"><h3><a href="{E(u)}" target="_blank">{E(u)}</a></h3><p>{badges}</p>{table}{fold}<div class="gal">{gal}</div></section>'

shared_html = ''.join(
    f'<div class="shared"><p>{badge(rid, len(v["pages"])).replace(" × ", " · pages: ")} in the site-wide <b>{E(region)}</b>. One template fix clears every page.</p>'
    f'<p class="q">{E(text)}</p><p>{E(BY_ID[rid]["fix"])}</p>'
    + (f'<a href="screenshots/{E(os.path.basename(v["hit"]["shot"]))}" target="_blank"><img loading="lazy" src="screenshots/{E(os.path.basename(v["hit"]["shot"]))}" alt="Screenshot"></a>' if v['hit'] else '')
    + '</div>' for (region, rid, text), v in shared.items()) or '<p class="muted">No forbidden wording in the site-wide header or footer.</p>'
exempt_rows = ''.join(f'<tr><td>{E(r["url"])}</td><td>{E(r["where"])}</td><td>{E(lbl(r["rule"]))}</td><td>{E(r["context"])}</td></tr>' for r in rows if r['exempt'])
clean = [p.get('url', p['req']) for p in pages if p.get('url', p['req']) not in page_rows]
rules_li = ''.join(f'<li><b>{E(r["severity"])}:</b> {E(r["label"])}{" (only on " + ", ".join(r["onlyPaths"]) + ")" if r.get("onlyPaths") else ""}. Fix: {E(r["fix"])}'
                   f'{"" if r.get("screenshot", True) else " Listed per page under a fold, not screenshotted."}</li>' for r in sorted(RULES, key=lambda r: ORDER.index(r['id'])))
in_sm = sum(1 for p in pages if p.get('inSitemap'))

doc = f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{E(a.site)} Wording Audit</title><style>
:root{{--bg:#fff;--fg:#1a1a1a;--mut:#666;--line:#e3e3e3;--card:#f7f5f5;--acc:#9b0d1f;--crit:#c62828;--hi:#ef6c00;--lo:#8d8d8d;--mark:#ffeb3b}}
@media (prefers-color-scheme:dark){{:root:not([data-theme="light"]){{--bg:#151314;--fg:#eee;--mut:#aaa;--line:#333;--card:#211d1e;--acc:#ff5a6e}}}}
:root[data-theme="dark"]{{--bg:#151314;--fg:#eee;--mut:#aaa;--line:#333;--card:#211d1e;--acc:#ff5a6e}}
body{{margin:0;background:var(--bg);color:var(--fg);font:15px/1.55 system-ui,-apple-system,Segoe UI,sans-serif}}
.wrap{{max-width:1200px;margin:0 auto;padding:24px 16px 80px}} h1{{margin:.2em 0;font-size:28px}} h2{{margin-top:48px;border-bottom:2px solid var(--acc);padding-bottom:6px}}
h3{{font-size:17px;word-break:break-all;margin:0 0 6px}} a{{color:var(--acc)}} .muted,small{{color:var(--mut)}}
.kpis{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin:18px 0}}
.kpi{{background:var(--card);border-radius:10px;padding:12px}} .kpi b{{display:block;font-size:26px}} .kpi span{{color:var(--mut);font-size:13px}}
.b{{display:inline-block;border-radius:999px;padding:1px 9px;font-size:12px;color:#fff;margin:2px 0;white-space:nowrap}} .crit{{background:var(--crit)}} .hi{{background:var(--hi)}} .lo{{background:var(--lo)}}
section,.shared{{background:var(--card);border-radius:12px;padding:16px;margin:16px 0;overflow:hidden}}
.tw{{overflow-x:auto}} table{{border-collapse:collapse;width:100%;font-size:13.5px;margin-top:8px}} th,td{{border-top:1px solid var(--line);padding:7px;text-align:left;vertical-align:top}} th{{font-size:12px;color:var(--mut)}}
mark{{background:var(--mark);color:#000;font-weight:700;padding:0 2px;outline:1px solid #e00}}
.q{{max-width:460px}} .gal{{display:grid;grid-template-columns:repeat(auto-fill,minmax(260px,1fr));gap:10px;margin-top:12px}}
.gal a{{display:block;text-decoration:none;font-size:12px;color:var(--mut)}} .gal img,.shared img{{width:100%;border-radius:6px;border:1px solid var(--line)}} .shared img{{max-width:640px}}
ul.toc{{columns:2;column-gap:28px;padding-left:18px}} ul.toc li{{break-inside:avoid;margin-bottom:4px;font-size:13.5px}} ol.dl li{{margin-bottom:6px;font-size:13px}}
details summary{{cursor:pointer;margin-top:10px;color:var(--mut)}} .rules li{{margin-bottom:4px}}
@media (max-width:700px){{ul.toc{{columns:1}} table{{font-size:12.5px}}}}
</style></head><body><div class="wrap">
<p class="muted">{E(re.sub(r"^https?://", "", pages[0]["url"]).split("/")[0] if pages else "")} · live check with Chrome DevTools MCP · {E(a.date)}</p>
<h1>{E(a.site)}: forbidden wording on live pages</h1>
<p>Every page in the live sitemap ({in_sm} URLs), plus {len(pages) - in_sm} internal pages linked from them, was opened in Chrome and its full rendered text checked against <b>{E(RULESET.get("name", "the house rules"))}</b>. Each finding names the page, the section heading it sits under, the sentence, the fix, and a screenshot with the word highlighted in <mark>yellow</mark> and its text block boxed in red.</p>
<div class="kpis"><div class="kpi"><b>{len(pages)}</b><span>pages checked</span></div><div class="kpi"><b>{len(word_pages)}</b><span>pages with forbidden wording</span></div>{kpis}</div>
<h2>Rules checked</h2><ul class="rules">{rules_li}
<li>Collapsed FAQ answers and accordions were opened for the screenshot; the table marks them "Collapsed".</li>
<li>Exempt: {E(RULESET.get("slugExemptNote", "none"))} Those are listed at the end, not as findings.</li></ul>
<h2>Pages to fix ({len(page_rows)})</h2><p class="muted">Ordered by {E(lbl(top))} count, then by number of screenshot-level findings.</p><ul class="toc">{toc}</ul>
<h2>Site-wide header and footer</h2>{shared_html}
<h2>Findings by page</h2>{sections}
<h2>Exempt text</h2><div class="tw"><table><thead><tr><th>Page</th><th>Field</th><th>Issue</th><th>Text</th></tr></thead><tbody>{exempt_rows or '<tr><td colspan=4>None</td></tr>'}</tbody></table></div>
<h2>Clean pages ({len(clean)})</h2><p class="muted">{E(", ".join(re.sub(r"^https?://[^/]+", "", c) or "/" for c in clean)) or "None"}</p>
<h2>Pages that served a not-found body ({len(broken)})</h2><p class="muted">Still broken after a retry. Not a wording issue, but worth a look if they are in the sitemap.</p><ul>{''.join(f'<li>{E(p["req"])}</li>' for p in broken) or '<li>None</li>'}</ul>
<h2>Pages that could not be checked ({len(errors)})</h2><ul>{''.join(f'<li>{E(p["req"])}: {E(p["error"])}</li>' for p in errors) or '<li>None</li>'}</ul>
<p class="muted">Method: headless Chrome at 1366×900 driven through Chrome DevTools MCP. Text was read from the rendered DOM, including collapsed FAQ text. Fixed pop-ups and chat widgets were hidden for the screenshots. The full list, with a Done column, is in {E(FILE_STEM)}_Forbidden_Wording_Findings.xlsx.</p>
</div></body></html>'''
open(os.path.join(a.out, f'{FILE_STEM}_Forbidden_Wording_Report.html'), 'w').write(doc)
zip_path = shutil.make_archive(a.out.rstrip('/'), 'zip', os.path.dirname(os.path.abspath(a.out)), os.path.basename(a.out.rstrip('/')))

print(json.dumps({'pages_checked': len(pages), 'pages_with_findings': len(page_rows), 'pages_with_forbidden_words': len(word_pages),
                  'by_rule': {lbl(k): v for k, v in by_rule.items()}, 'site_wide': len(shared), 'exempt': sum(r['exempt'] for r in rows),
                  'not_found_after_retry': [p['req'] for p in broken], 'errors': [p['req'] for p in errors],
                  'screenshots': len(used), 'zip': zip_path, 'zip_mb': round(os.path.getsize(zip_path) / 1e6, 1)}, indent=1))
